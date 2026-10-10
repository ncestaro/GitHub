"""Crea il PowerPoint per il teatro con video e fermi immagine a tutto schermo.

Struttura (nessuna transizione grafica, il passaggio è un taglio netto):
  1  video poltrona      parte da solo, alla fine passa da solo alla slide 2
  2  fermo poltrona      resta ferma fino al clic
  3  video contatore     parte da solo, alla fine passa da solo alla slide 4
  4  fermo importo       resta ferma fino al clic
  5  video 8 poltrone    parte da solo, alla fine passa da solo alla slide 6
  6  fermo 8 poltrone    resta ferma (fine)

Uso: python3 tools/crea_pptx.py [--formato 169|43]
Legge i file da export/ (16:9) o export/4_3/ (4:3), scrive il .pptx nella stessa cartella.
"""
import argparse
import pathlib
import subprocess
import tempfile

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu

here = pathlib.Path(__file__).resolve().parent.parent
ap = argparse.ArgumentParser()
ap.add_argument("--formato", default="169", choices=["169", "43"])
args = ap.parse_args()

src = here / "export" if args.formato == "169" else here / "export" / "4_3"
W = Emu(12192000 if args.formato == "169" else 9144000)   # 13,33" o 10" di larghezza
H = Emu(6858000)                                          # 7,5" di altezza
out = src / f"Raccolta_9320_{'16-9' if args.formato == '169' else '4-3'}.pptx"

SLIDES = [
    ("video", "1_poltrona.mp4"), ("image", "2_fermo_poltrona.png"),
    ("video", "3_contatore.mp4"), ("image", "4_fermo_importo.png"),
    ("video", "5_poltrone.mp4"), ("image", "6_fermo_poltrone.png"),
]

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
TIMING = """<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="{ms}" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn>
<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
</p:seq>
<p:video><p:cMediaNode vol="80000"><p:cTn id="7" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"""


def duration_ms(path):
    r = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(path)],
                       capture_output=True, text=True, check=True)
    return round(float(r.stdout) * 1000)


def set_transition(sld, adv_ms=None):
    """Nessun effetto di transizione; con adv_ms la slide avanza da sola dopo adv_ms."""
    for old in sld.findall(qn("p:transition")):
        sld.remove(old)
    tr = etree.SubElement(sld, qn("p:transition"))
    if adv_ms is not None:
        tr.set("advTm", str(adv_ms))
    # ordine richiesto: cSld, clrMapOvr, transition, timing, extLst
    anchor = sld.find(qn("p:clrMapOvr"))
    anchor.addnext(tr)


prs = Presentation()
prs.slide_width, prs.slide_height = W, H
blank = prs.slide_layouts[6]
with tempfile.TemporaryDirectory() as tmp:
    for i, (kind, name) in enumerate(SLIDES, 1):
        path = src / name
        slide = prs.slides.add_slide(blank)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor(0xF0, 0xDF, 0xCB)
        sld = slide._element
        if kind == "image":
            slide.shapes.add_picture(str(path), 0, 0, W, H)
            set_transition(sld)
        else:
            poster = pathlib.Path(tmp) / f"poster{i}.png"
            subprocess.run(["ffmpeg", "-v", "error", "-y", "-i", str(path), "-frames:v", "1", str(poster)], check=True)
            movie = slide.shapes.add_movie(str(path), 0, 0, W, H, poster_frame_image=str(poster), mime_type="video/mp4")
            movie.name = name
            ms = duration_ms(path)
            for old in sld.findall(qn("p:timing")):
                sld.remove(old)
            timing = etree.fromstring(TIMING.format(P=P, ms=ms, spid=movie.shape_id).replace("\n", ""))
            sld.find(qn("p:clrMapOvr")).addnext(timing)
            set_transition(sld, ms + 100)   # passa alla slide del fermo immagine appena finito il video

prs.core_properties.title = "Raccolta fondi – Storie di donne straordinariamente semplici"
prs.save(out)
print("ok", out.relative_to(here))
